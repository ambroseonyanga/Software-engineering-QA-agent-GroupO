from pathlib import Path
from docx import Document
import sys


def convert_docx_to_markdown(input_file, output_file):
    document = Document(input_file)

    lines = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        style = paragraph.style.name.lower() if paragraph.style else ""

        if "title" in style:
            lines.append(f"# {text}")
        elif "heading 1" in style:
            lines.append(f"## {text}")
        elif "heading 2" in style:
            lines.append(f"### {text}")
        else:
            lines.append(text)

        lines.append("")

    for table in document.tables:
        if not table.rows:
            continue

        rows = []

        for row in table.rows:
            cells = [
                cell.text.replace("\n", " ").strip()
                for cell in row.cells
            ]
            rows.append(cells)

        if rows:
            header = rows[0]

            lines.append("| " + " | ".join(header) + " |")
            lines.append(
                "| " + " | ".join(["---"] * len(header)) + " |"
            )

            for row in rows[1:]:
                lines.append("| " + " | ".join(row) + " |")

            lines.append("")

    output_file = Path(output_file)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    output_file.write_text(
        "\n".join(lines),
        encoding="utf-8"
    )

    print(f"Converted: {input_file}")
    print(f"Created:   {output_file}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(
            "Usage: python docx_to_md.py "
            "<input.docx> <output.md>"
        )
        sys.exit(1)

    convert_docx_to_markdown(
        sys.argv[1],
        sys.argv[2]
    )
