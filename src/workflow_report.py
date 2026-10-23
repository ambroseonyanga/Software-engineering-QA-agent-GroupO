"""Editable reports derived from recorded execution evidence, not model verdicts."""
from collections import Counter
from io import BytesIO
import re

try:
    from .workflow_service import uganda_time
except ImportError:
    from workflow_service import uganda_time

from docx import Document
from docx.shared import Inches, Pt, RGBColor


def clean(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value))


def report_bytes(state):
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    document.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
    document.add_heading(clean(state['name']) + ' QA report', 0)
    counts = Counter(result['status'] for result in state['results'])
    approved = state.get('approval') or {'tests': {}, 'at': 'Not approved', 'reviewer': 'None'}
    document.add_paragraph(
        f"This report records the tests selected for {clean(state['name'])} and their observed sandbox results. "
        f"{counts['passed']} passed and {counts['failed']} failed. "
        'Execution errors, skipped tests and tests that did not run are listed separately below. '
        'These results cover the selected tests, not the whole application or SRS.')
    document.add_heading('Run summary', level=1)
    for label, value in [('Requirements document', state['pdf_name']), ('Run state', state['status']),
                         ('Session', state['id']), ('Finished EAT (UTC+3)', uganda_time(state.get('finished') or 'Not finished')),
                         ('Approved by', approved['reviewer']), ('Approved EAT (UTC+3)', uganda_time(approved['at'])),
                         ('Environment', 'Django with temporary SQLite' if state['adapter'] == 'budgetme' else 'Python unittest'),
                         ('Suggested tests', len(state['tests'])), ('Approved tests', len(approved['tests']))]:
        document.add_paragraph(f'{label}: {clean(value)}')
    if state.get('error'):
        document.add_paragraph('Run problem: ' + clean(state['error']))
    document.add_heading('Results', level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = 'Light Shading Accent 1'
    for cell, label in zip(table.rows[0].cells, ['Test', 'Requirement', 'Approved', 'Result']):
        cell.text = label
    results = {result['id']: result for result in state['results']}
    for test in state['tests']:
        result = results.get(test['id'])
        values = [test['id'], test['requirement'], 'Yes' if test['id'] in approved['tests'] else 'No',
                  result['status'] if result else 'Not run']
        for cell, value in zip(table.add_row().cells, values):
            cell.text = clean(value)
    for test in state['tests']:
        document.add_heading(clean(test['id'] + ' ' + test['title']), level=2)
        document.add_paragraph('Expected: ' + clean(test['expected']))
        document.add_paragraph('Coverage limit: ' + clean(test['limitations']))
        result = results.get(test['id'])
        if result:
            document.add_paragraph(f"Observed: {result['status']}. Duration: {result['seconds']} seconds. Exit code: {result['exit_code']}.")
            detail = result['detail'] or ('Assertions completed successfully.' if result['status'] == 'passed' else 'See execution output.')
            document.add_paragraph(clean(detail[:3500]))
            if result['status'] != 'passed':
                paragraph = document.add_paragraph(clean(result['output'][-3500:]))
                for run in paragraph.runs:
                    run.font.name = 'Consolas'
                    run.font.size = Pt(8)
        else:
            document.add_paragraph('Not run. ' + ('Approved, but execution did not finish.' if test['id'] in approved['tests'] else 'Not selected for approval.'))
    document.add_heading('Scope and next steps', level=1)
    document.add_paragraph('Suggestions were generated from selected PDF and code excerpts. Review each failed test before treating it as an application defect. Setup errors require a working test environment before the requirement can be evaluated. Scanned PDF OCR and browser automation are outside this prototype.')
    if state['adapter'] == 'budgetme':
        document.add_paragraph('External integrations were disabled. SQLite results do not establish Supabase permissions, PostgreSQL concurrency behaviour or live OAuth behaviour.')
    document.add_paragraph('All executions used isolated containers with no network, no host folder mounts, a read-only root filesystem and resource limits. Each proposed test had a 90-second time limit.')
    document.add_heading('Evidence references', level=1)
    for label, value in [('Requirements input SHA256', state['pdf_hash']), ('Source snapshot SHA256', state['snapshot_hash']),
                         ('Docker image', state.get('image_id', 'Image preparation did not complete')),
                         ('Generation model', state.get('model', 'Not recorded'))]:
        document.add_paragraph(label + ': ' + clean(value))
    for test in state['tests']:
        document.add_paragraph(test['id'] + ' code SHA256: ' + test.get('hash', 'No code generated; test not selected.'))
    document.core_properties.title = clean(state['name']) + ' QA report'
    document.core_properties.author = 'QA Workspace'
    output = BytesIO()
    document.save(output)
    return output.getvalue()
